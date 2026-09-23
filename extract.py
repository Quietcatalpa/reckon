"""读取“容易读”的文件内容片段，交给 Laya 参考。读不了就返回 None。"""
import json
import os
import re
import zipfile

TEXT_EXTS = set(
    ".txt .md .log .csv .tsv .json .xml .yaml .yml .ini .cfg .toml .tex .bib .py .js .ts .jsx .tsx "
    ".java .c .cpp .h .hpp .m .r .go .rs .cs .html .css .sql .sh .bat .ps1".split()
)
READABLE_EXTS = TEXT_EXTS | {".docx", ".pptx", ".xlsx", ".xlsm", ".pdf", ".zip", ".ipynb"}
_WS = re.compile(r"\s+")


def _clean(s, limit):
    return _WS.sub(" ", s).strip()[:limit] or None


def _text(path, limit):
    with open(path, "rb") as f:
        raw = f.read(limit * 4)
    if b"\x00" in raw[:1024]:
        return None  # 其实是二进制
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", "replace")


def _docx(path, limit):
    from docx import Document
    out, n = [], 0
    for para in Document(path).paragraphs:
        if para.text.strip():
            out.append(para.text)
            n += len(para.text)
            if n > limit:
                break
    return "\n".join(out)


def _pptx(path, limit):
    from pptx import Presentation
    out, n = [], 0
    for i, slide in enumerate(Presentation(path).slides):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
                n += len(out[-1])
        if n > limit or i > 10:
            break
    return "\n".join(out)


def _xlsx(path, limit):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = ["工作表：" + "、".join(wb.sheetnames[:10])]
        for row in wb.worksheets[0].iter_rows(max_row=12, values_only=True):
            cells = [str(v) for v in row[:12] if v is not None]
            if cells:
                out.append(" | ".join(cells))
        return "\n".join(out)
    finally:
        wb.close()


def _pdf(path, limit):
    import fitz
    with fitz.open(path) as doc:
        out = []
        title = (doc.metadata or {}).get("title")
        if title:
            out.append("标题：" + title)
        out.append(f"共 {doc.page_count} 页")
        for page in doc.pages(0, min(2, doc.page_count)):
            out.append(page.get_text())
        return "\n".join(out)


def _zip(path, limit):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return f"压缩包内共 {len(names)} 个文件：" + "；".join(names[:40])


def _ipynb(path, limit):
    with open(path, encoding="utf-8") as f:
        nb = json.load(f)
    return "\n".join("".join(c.get("source", "")) for c in nb.get("cells", [])[:8])


_READERS = {".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx,
            ".pdf": _pdf, ".zip": _zip, ".ipynb": _ipynb}


OFFICE_MAX_BYTES = 100 << 20  # 超大的 docx/pptx/xlsx 解析很慢（整包读进内存），只看文件信息
OFFICE_EXTS = {".docx", ".pptx", ".xlsx", ".xlsm"}


def snippet(path, ext, limit=600):
    if ext not in READABLE_EXTS:
        return None
    reader = _READERS.get(ext, _text)
    try:
        if ext in OFFICE_EXTS and os.path.getsize(path) > OFFICE_MAX_BYTES:
            return None
        return _clean(reader(path, limit), limit)
    except Exception:
        return None
